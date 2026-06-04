"""Build Excel for manual eye-balling of all detected acronym expansions.

Sheets:
  1_All_Acronyms_Aggregated  - master list per acronym (unique)
  2_Per_Paper_Detection      - tiap deteksi di tiap paper (raw, ribuan rows)
  3_Frequency_Analysis       - acronym paling sering muncul lintas paper
  4_Suspect_Review           - kandidat false positive (multi-meaning, weird length, etc)
  5_Curated_Dict             - 175 acronym dari Layer 2 curated dict
"""
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from acronym_expander_sh import extract_pairs_schwartz_hearst
from acronym_dict import COMMON_MEDICAL_ACRONYMS, AMBIGUOUS_ACRONYMS

from datasets import load_dataset

HERE = Path(__file__).parent
OUT = HERE / 'acronym_audit.xlsx'


# ============================================================
# Style
# ============================================================
NAVY = '1F4E79'
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor=NAVY)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
BODY_FONT = Font(name='Calibri', size=10)
WRAP = Alignment(vertical='top', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN = Side(border_style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LRED = 'FEE2E2'; LGRN = 'D1FAE5'; LYEL = 'FEF3C7'
LBLU = 'DBEAFE'; LGRY = 'F3F4F6'; LPUR = 'E9D5FF'


def style_header(ws, row=1, last_col=None):
    if last_col is None:
        last_col = ws.max_column
    for col in range(1, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN
        c.border = BORDER
    ws.row_dimensions[row].height = 32


# ============================================================
# Extract context snippet for an acronym in a paper
# ============================================================
def find_context_snippet(text, acronym, window=80):
    """Find first occurrence of '(ACRONYM)' and return snippet around it."""
    pattern = rf'\b\w[^.;()]*\(\s*{re.escape(acronym)}\s*\)'
    m = re.search(pattern, text)
    if m:
        start = max(0, m.start() - 10)
        end = min(len(text), m.end() + 30)
        return text[start:end].strip()
    # Fallback: just find acronym
    m = re.search(rf'\b{re.escape(acronym)}\b', text)
    if m:
        start = max(0, m.start() - window)
        end = min(len(text), m.end() + window)
        return text[start:end].strip()
    return ''


def main():
    print('Loading PubMedQA dataset (first 500 papers)...')
    ds = load_dataset('qiaojin/PubMedQA', 'pqa_labeled')['train'].select(range(500))
    print(f'  {len(ds)} papers loaded')

    # Per-paper detection
    print('\nExtracting acronyms via Schwartz-Hearst from each paper...')
    per_paper_detections = []  # list of dicts
    paper_acronym_count = defaultdict(set)  # acronym -> set of pubids
    aggregate_expansions = defaultdict(Counter)  # acronym -> Counter of expansions

    for i, item in enumerate(ds):
        pubid = str(item['pubid'])
        question = item['question']
        sections = item['context']['contexts']
        labels = item['context']['labels']
        full_text = ' '.join(sections)

        # Schwartz-Hearst Layer 1 only
        sh_pairs = extract_pairs_schwartz_hearst(full_text)
        for ac, exp in sh_pairs.items():
            snippet = find_context_snippet(full_text, ac)
            per_paper_detections.append({
                'pubid': pubid,
                'question': question[:80],
                'acronym': ac,
                'expansion': exp,
                'layer': 'Layer 1 (in-paper)',
                'snippet': snippet[:200],
            })
            paper_acronym_count[ac].add(pubid)
            aggregate_expansions[ac][exp] += 1

        if (i + 1) % 100 == 0:
            print(f'  Processed {i+1}/{len(ds)}')

    print(f'\nTotal in-paper detections: {len(per_paper_detections)}')
    print(f'Unique acronyms (Layer 1): {len(paper_acronym_count)}')

    # ============================================================
    # Build Excel
    # ============================================================
    wb = Workbook()
    wb.remove(wb.active)

    # ----- Sheet 1: Aggregated per acronym -----
    ws1 = wb.create_sheet('1_All_Acronyms_Aggregated')
    ws1['A1'] = ('All Acronyms Detected by Schwartz-Hearst (Layer 1: in-paper extraction). '
                 f'Total unique acronyms: {len(paper_acronym_count)}.')
    ws1['A1'].font = Font(size=12, bold=True, color=NAVY)
    ws1.merge_cells('A1:F1')

    ws1['A2'] = ('Sortable list of all detected acronyms. "n_papers" = berapa banyak paper '
                 'yang punya acronym ini. "Most common expansion" = expansion yang paling sering '
                 'muncul di lintas paper. "All expansions" = semua expansion yang pernah muncul '
                 '(format: expansion (count)).')
    ws1['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws1['A2'].alignment = WRAP
    ws1.merge_cells('A2:F2')
    ws1.row_dimensions[2].height = 36

    headers = ['Acronym', 'n_papers', 'Most common expansion', 'All expansions (count)',
               'In curated dict?', 'Status check']
    for i, h in enumerate(headers, 1):
        ws1.cell(row=4, column=i, value=h)
    style_header(ws1, row=4, last_col=len(headers))

    # Sort by frequency
    sorted_acronyms = sorted(paper_acronym_count.items(), key=lambda x: -len(x[1]))

    for row_i, (ac, papers) in enumerate(sorted_acronyms, 5):
        n_papers = len(papers)
        expansions = aggregate_expansions[ac]
        most_common = expansions.most_common(1)[0][0]
        all_exp = '; '.join(f'{e} ({c})' for e, c in expansions.most_common())
        in_dict = 'YES' if ac in COMMON_MEDICAL_ACRONYMS else 'NO'

        # Status check heuristik
        status = ''
        if len(expansions) > 1:
            status = '⚠ Multiple expansions (review)'
        elif ac in AMBIGUOUS_ACRONYMS:
            status = '⚠ Listed as ambiguous'
        elif len(most_common) < 5:
            status = '⚠ Suspiciously short expansion'
        elif most_common.startswith(('and ', 'or ', 'in ', 'of ', 'the ', 'a ')):
            status = '⚠ Starts with stopword'
        else:
            status = 'OK'

        cells = [ac, n_papers, most_common, all_exp, in_dict, status]
        for col_i, v in enumerate(cells, 1):
            c = ws1.cell(row=row_i, column=col_i, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP if col_i in (3, 4) else Alignment(horizontal='center' if col_i in (2, 5) else 'left', vertical='top', wrap_text=True)
            c.border = BORDER

        # Color status
        if 'OK' in status:
            ws1.cell(row=row_i, column=6).fill = PatternFill('solid', fgColor=LGRN)
        elif '⚠' in status:
            ws1.cell(row=row_i, column=6).fill = PatternFill('solid', fgColor=LYEL)

    widths = [12, 10, 38, 50, 14, 32]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = 'A5'
    ws1.auto_filter.ref = f'A4:F{4 + len(sorted_acronyms)}'

    # ----- Sheet 2: Per-paper detection (raw) -----
    ws2 = wb.create_sheet('2_Per_Paper_Detection')
    ws2['A1'] = ('Per-Paper Acronym Detection — RAW. '
                 f'Total detections: {len(per_paper_detections)}. '
                 'Setiap baris = 1 acronym terdetect di 1 paper.')
    ws2['A1'].font = Font(size=12, bold=True, color=NAVY)
    ws2.merge_cells('A1:F1')

    ws2['A2'] = ('Filter via auto-filter untuk fokus per acronym/pubid. '
                 'Snippet berisi konteks 200 karakter sekitar acronym definition '
                 '(untuk verifikasi manual).')
    ws2['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws2.merge_cells('A2:F2')

    headers = ['PMID', 'Acronym', 'Expansion', 'Layer', 'Question (preview)', 'Context Snippet']
    for i, h in enumerate(headers, 1):
        ws2.cell(row=4, column=i, value=h)
    style_header(ws2, row=4, last_col=len(headers))

    for row_i, det in enumerate(per_paper_detections, 5):
        cells = [det['pubid'], det['acronym'], det['expansion'], det['layer'],
                 det['question'], det['snippet']]
        for col_i, v in enumerate(cells, 1):
            c = ws2.cell(row=row_i, column=col_i, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
            c.border = BORDER
        # PMID hyperlink
        ws2.cell(row=row_i, column=1).hyperlink = f'https://pubmed.ncbi.nlm.nih.gov/{det["pubid"]}/'
        ws2.cell(row=row_i, column=1).font = Font(color=NAVY, underline='single', size=10)

    widths = [12, 12, 35, 22, 40, 60]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = 'A5'
    ws2.auto_filter.ref = f'A4:F{4 + len(per_paper_detections)}'

    # ----- Sheet 3: Frequency analysis -----
    ws3 = wb.create_sheet('3_Frequency_Analysis')
    ws3['A1'] = 'Top 50 Acronyms by Frequency'
    ws3['A1'].font = Font(size=12, bold=True, color=NAVY)
    ws3.merge_cells('A1:E1')

    ws3['A2'] = ('Acronym yang muncul di banyak paper. Top entries kemungkinan '
                 'common medical terms dengan multiple meanings cross-paper.')
    ws3['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws3.merge_cells('A2:E2')

    headers = ['Rank', 'Acronym', 'n_papers', 'Top expansion', 'Variant count']
    for i, h in enumerate(headers, 1):
        ws3.cell(row=4, column=i, value=h)
    style_header(ws3, row=4, last_col=len(headers))

    for row_i, (ac, papers) in enumerate(sorted_acronyms[:50], 5):
        n_papers = len(papers)
        expansions = aggregate_expansions[ac]
        most_common = expansions.most_common(1)[0][0]
        cells = [row_i - 4, ac, n_papers, most_common, len(expansions)]
        for col_i, v in enumerate(cells, 1):
            c = ws3.cell(row=row_i, column=col_i, value=v)
            c.font = BODY_FONT
            c.alignment = CENTER if col_i != 4 else Alignment(vertical='center', wrap_text=True)
            c.border = BORDER
        if len(expansions) > 1:
            ws3.cell(row=row_i, column=5).fill = PatternFill('solid', fgColor=LYEL)
            ws3.cell(row=row_i, column=5).font = Font(bold=True, color='B45309')

    widths = [8, 12, 10, 45, 14]
    for i, w in enumerate(widths, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ----- Sheet 4: Suspect/review -----
    ws4 = wb.create_sheet('4_Suspect_Review')
    ws4['A1'] = 'Acronyms yang Perlu Direview Manual'
    ws4['A1'].font = Font(size=12, bold=True, color=NAVY)
    ws4.merge_cells('A1:G1')

    ws4['A2'] = ('Filter heuristic: (a) multiple expansions di lintas paper, '
                 '(b) listed as ambiguous, (c) expansion suspiciously short/odd, '
                 '(d) expansion mulai stopword. Eyeball masing-masing.')
    ws4['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws4.merge_cells('A2:G2')
    ws4.row_dimensions[2].height = 32

    headers = ['Acronym', 'n_papers', 'n_variants', 'All expansions', 'Reason flagged', 'Sample PMID', 'Action recommendation']
    for i, h in enumerate(headers, 1):
        ws4.cell(row=4, column=i, value=h)
    style_header(ws4, row=4, last_col=len(headers))

    suspect_count = 0
    for ac, papers in sorted_acronyms:
        expansions = aggregate_expansions[ac]
        n_variants = len(expansions)
        most_common = expansions.most_common(1)[0][0]

        flags = []
        if n_variants > 1:
            flags.append(f'{n_variants} variant expansions')
        if ac in AMBIGUOUS_ACRONYMS:
            flags.append('listed as ambiguous')
        if len(most_common) < 5:
            flags.append('expansion <5 chars')
        if any(most_common.startswith(w + ' ') for w in ['and', 'or', 'in', 'of', 'the', 'a', 'an']):
            flags.append('starts with stopword')
        if len(ac) <= 2 and most_common in ['and', 'or', 'or']:
            flags.append('common word as expansion')

        if not flags:
            continue

        suspect_count += 1
        all_exp = '; '.join(f'{e} ({c})' for e, c in expansions.most_common(5))
        sample_pmid = next(iter(papers))

        # Action recommendation
        if 'ambiguous' in ' '.join(flags):
            action = 'Override with curated dict per paper'
        elif n_variants > 3:
            action = 'Add to manual override list'
        elif 'short' in ' '.join(flags):
            action = 'Validate: likely false positive'
        elif 'stopword' in ' '.join(flags):
            action = 'Improve extractor: strip leading stopwords'
        else:
            action = 'Keep but monitor'

        cells = [ac, len(papers), n_variants, all_exp,
                 ', '.join(flags), sample_pmid, action]
        for col_i, v in enumerate(cells, 1):
            c = ws4.cell(row=4 + suspect_count, column=col_i, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
            c.border = BORDER
        # PMID link
        ws4.cell(row=4 + suspect_count, column=6).hyperlink = f'https://pubmed.ncbi.nlm.nih.gov/{sample_pmid}/'
        ws4.cell(row=4 + suspect_count, column=6).font = Font(color=NAVY, underline='single', size=10)
        # Color severity
        if n_variants > 3:
            ws4.cell(row=4 + suspect_count, column=3).fill = PatternFill('solid', fgColor=LRED)
        elif n_variants > 1:
            ws4.cell(row=4 + suspect_count, column=3).fill = PatternFill('solid', fgColor=LYEL)

    widths = [12, 10, 12, 50, 30, 14, 38]
    for i, w in enumerate(widths, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = 'A5'
    if suspect_count > 0:
        ws4.auto_filter.ref = f'A4:G{4 + suspect_count}'

    # ----- Sheet 5: Curated dict -----
    ws5 = wb.create_sheet('5_Curated_Dict')
    ws5['A1'] = 'Layer 2 Curated Medical Acronym Dictionary'
    ws5['A1'].font = Font(size=12, bold=True, color=NAVY)
    ws5.merge_cells('A1:D1')

    ws5['A2'] = (f'Total {len(COMMON_MEDICAL_ACRONYMS)} acronym yang di-hardcode sebagai fallback '
                 'kalau in-paper detection gagal. Termasuk yang paling common di biomedical text.')
    ws5['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws5.merge_cells('A2:D2')

    headers = ['Acronym', 'Expansion', 'Listed as ambiguous?', 'Detected in-paper di berapa paper?']
    for i, h in enumerate(headers, 1):
        ws5.cell(row=4, column=i, value=h)
    style_header(ws5, row=4, last_col=len(headers))

    for row_i, (ac, exp) in enumerate(sorted(COMMON_MEDICAL_ACRONYMS.items()), 5):
        is_ambig = 'YES' if ac in AMBIGUOUS_ACRONYMS else 'NO'
        n_in_paper = len(paper_acronym_count.get(ac, set()))
        cells = [ac, exp, is_ambig, n_in_paper]
        for col_i, v in enumerate(cells, 1):
            c = ws5.cell(row=row_i, column=col_i, value=v)
            c.font = BODY_FONT
            c.alignment = CENTER if col_i in (1, 3, 4) else Alignment(vertical='center', wrap_text=True)
            c.border = BORDER
        if is_ambig == 'YES':
            ws5.cell(row=row_i, column=3).fill = PatternFill('solid', fgColor=LYEL)

    widths = [14, 50, 18, 32]
    for i, w in enumerate(widths, 1):
        ws5.column_dimensions[get_column_letter(i)].width = w
    ws5.freeze_panes = 'A5'
    ws5.auto_filter.ref = f'A4:D{4 + len(COMMON_MEDICAL_ACRONYMS)}'

    # Save
    wb.save(OUT)
    print(f'\n[Saved] {OUT}')
    print(f'  Size: {OUT.stat().st_size / 1024:.1f} KB')
    print(f'\nSheets:')
    print(f'  1. All Acronyms Aggregated  : {len(sorted_acronyms)} unique acronyms')
    print(f'  2. Per-Paper Detection      : {len(per_paper_detections)} raw detections')
    print(f'  3. Frequency Analysis        : top 50')
    print(f'  4. Suspect Review            : {suspect_count} flagged for manual review')
    print(f'  5. Curated Dict              : {len(COMMON_MEDICAL_ACRONYMS)} fallback entries')


if __name__ == '__main__':
    main()
