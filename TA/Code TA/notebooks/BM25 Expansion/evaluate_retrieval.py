"""Compute 6 retrieval metrics + per-section retrieval rate for 3 versions
(Original, V1 Naive Expansion, V2 Schwartz-Hearst).

Source-paper-based relevance: a chunk is relevant if its pubid matches the
source paper of the query (per PubMedQA construction by Jin et al. 2019).

Outputs:
  - retrieval_metrics.json  (machine-readable, for HTML embed)
  - retrieval_metrics_minimal.xlsx (Excel summary + per-question matrix)
"""
import json
import math
import pickle
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@dataclass
class Document:
    text: str; pubid: str; question: str
    section_label: str; answer: str; decision: str


import __main__
__main__.Document = Document


HERE = Path(__file__).parent
RESULTS = HERE.parent.parent / 'results'
RESULTS_BM25 = RESULTS / 'BM25_Expansion'
NOTEBOOKS = HERE.parent

K = 5  # top-k


# ============================================================
# Helpers
# ============================================================
def load_phase1(path):
    if not path.exists():
        return None
    with open(path, encoding='utf-8') as f:
        return json.load(f)['results']


def build_text_to_meta_map(documents):
    """Map text[:80] -> (pubid, section_label) for original baseline mapping."""
    return {d.text[:80]: (d.pubid, d.section_label) for d in documents}


def derive_pubids_sections(result, text_map):
    """Derive pubids and sections from contexts text via prefix matching."""
    pubids, sections = [], []
    for ctx in result['contexts']:
        meta = text_map.get(ctx[:80])
        if meta:
            pubids.append(meta[0])
            sections.append(meta[1])
        else:
            pubids.append('?')
            sections.append('?')
    return pubids, sections


# ============================================================
# Metric computation
# ============================================================
def compute_metrics(retrieved_pubids, retrieved_sections, source_pubid, n_source_sections):
    """Compute 6 retrieval metrics for a single query."""
    k = len(retrieved_pubids)
    rel = [1 if p == source_pubid else 0 for p in retrieved_pubids]
    n_rel = sum(rel)

    # Hit Rate@K
    hit = 1 if n_rel > 0 else 0

    # Precision@K
    precision = n_rel / k

    # Recall@K (capped at 1.0)
    recall = min(n_rel / max(n_source_sections, 1), 1.0)

    # MRR
    mrr = 0.0
    for i, r in enumerate(rel):
        if r == 1:
            mrr = 1.0 / (i + 1)
            break

    # nDCG@K (binary relevance)
    dcg = sum(r / math.log2(i + 2) for i, r in enumerate(rel))
    n_ideal = min(n_source_sections, k)
    idcg = sum(1 / math.log2(i + 2) for i in range(n_ideal))
    ndcg = dcg / idcg if idcg > 0 else 0.0

    # MAP@K
    if n_rel == 0:
        map_score = 0.0
    else:
        cum_rel = 0
        prec_at_rel = []
        for i, r in enumerate(rel):
            if r == 1:
                cum_rel += 1
                prec_at_rel.append(cum_rel / (i + 1))
        map_score = sum(prec_at_rel) / n_source_sections

    # Per-section retrieval (which source sections were retrieved)
    src_sections_retrieved = set()
    for p, s in zip(retrieved_pubids, retrieved_sections):
        if p == source_pubid:
            src_sections_retrieved.add(s)

    return {
        'hit@5': hit,
        'precision@5': precision,
        'recall@5': recall,
        'mrr': mrr,
        'ndcg@5': ndcg,
        'map@5': map_score,
        'n_relevant_in_topK': n_rel,
        'src_sections_retrieved': list(src_sections_retrieved),
    }


# ============================================================
# Aggregation
# ============================================================
def evaluate_version(results, version_label, text_map_for_derivation, source_meta):
    """Compute aggregate metrics for one version (Original/V1/V2).

    source_meta: dict {pubid: list of section_labels in corpus}
    """
    per_query = []
    for r in results:
        idx = r['idx']
        target_pubid = str(r['pubid'])

        # Get pubids and sections
        if 'context_pubids' in r and 'context_sections' in r:
            pubids = r['context_pubids']
            sections = r['context_sections']
        else:
            pubids, sections = derive_pubids_sections(r, text_map_for_derivation)

        n_source_sections = len(source_meta.get(target_pubid, []))

        m = compute_metrics(pubids, sections, target_pubid, n_source_sections)
        m['idx'] = idx
        m['source_pubid'] = target_pubid
        m['source_n_sections'] = n_source_sections
        m['retrieved_pubids'] = pubids
        m['retrieved_sections'] = sections
        per_query.append(m)

    # Aggregate
    n = len(per_query)
    agg = {
        'hit@5': sum(m['hit@5'] for m in per_query) / n,
        'precision@5': sum(m['precision@5'] for m in per_query) / n,
        'recall@5': sum(m['recall@5'] for m in per_query) / n,
        'mrr': sum(m['mrr'] for m in per_query) / n,
        'ndcg@5': sum(m['ndcg@5'] for m in per_query) / n,
        'map@5': sum(m['map@5'] for m in per_query) / n,
    }

    # Per-section retrieval rate (across all queries where source has that section)
    section_attempts = defaultdict(int)  # how many queries had this section in source
    section_hits = defaultdict(int)  # how many of those had it retrieved
    for m in per_query:
        target = m['source_pubid']
        avail_sections = source_meta.get(target, [])
        for sec in avail_sections:
            section_attempts[sec] += 1
        for sec in m['src_sections_retrieved']:
            if sec in avail_sections:  # validate it's a real section of source
                section_hits[sec] += 1

    per_section = {}
    for sec, n_attempts in section_attempts.items():
        per_section[sec] = {
            'attempts': n_attempts,
            'hits': section_hits[sec],
            'rate': section_hits[sec] / n_attempts if n_attempts > 0 else 0,
        }

    return {
        'label': version_label,
        'n_queries': n,
        'aggregate': agg,
        'per_section': per_section,
        'per_query': per_query,
    }


# ============================================================
# MAIN
# ============================================================
def main():
    print('Loading BM25 indexes (for source meta + text mapping)...')

    # Load original BM25 to get document map (for original baseline derivation)
    with open(NOTEBOOKS / 'pubmedqa_bm25.pkl', 'rb') as f:
        saved_orig = pickle.load(f)
    documents_orig = saved_orig['documents']
    text_map_orig = build_text_to_meta_map(documents_orig)
    print(f'  Original BM25 docs: {len(documents_orig)}')

    # Build source_meta: for each pubid, what sections are in corpus?
    source_meta = defaultdict(list)
    for d in documents_orig:
        source_meta[d.pubid].append(d.section_label)
    print(f'  Unique source papers in corpus: {len(source_meta)}')

    # Load 3 result sets
    print('\nLoading 3 result sets...')
    orig = load_phase1(RESULTS / 'baseline_openai_phase1_answers.json')
    v1 = load_phase1(RESULTS_BM25 / 'bm25expanded_baseline_openai_phase1_answers.json')
    v2 = load_phase1(RESULTS_BM25 / 'sh_baseline_openai_phase1_answers.json')

    print(f'  Original: {len(orig) if orig else "MISSING"} samples')
    print(f'  V1 Naive: {len(v1) if v1 else "MISSING"} samples')
    print(f'  V2 SH   : {len(v2) if v2 else "MISSING"} samples')

    if not (orig and v1 and v2):
        raise RuntimeError('Missing result files. Run baseline notebooks first.')

    # Evaluate all 3
    print('\nEvaluating Original...')
    eval_orig = evaluate_version(orig, 'Original', text_map_orig, source_meta)
    print('Evaluating V1 (Naive)...')
    eval_v1 = evaluate_version(v1, 'V1 Naive', text_map_orig, source_meta)
    print('Evaluating V2 (Schwartz-Hearst)...')
    eval_v2 = evaluate_version(v2, 'V2 Schwartz-Hearst', text_map_orig, source_meta)

    # Print summary
    print('\n' + '=' * 80)
    print('AGGREGATE METRICS')
    print('=' * 80)
    print(f'{"Metric":<18} {"Original":>12} {"V1 Naive":>12} {"V2 SH":>12} {"Δ V2-Orig":>12}')
    print('-' * 80)
    for metric in ['hit@5', 'precision@5', 'recall@5', 'mrr', 'ndcg@5', 'map@5']:
        o = eval_orig['aggregate'][metric]
        v = eval_v1['aggregate'][metric]
        s = eval_v2['aggregate'][metric]
        d = (s - o) * 100
        print(f'  {metric:<16} {o:>11.4f}  {v:>11.4f}  {s:>11.4f}  {d:+11.2f} pp')

    # Per-section breakdown
    print('\n' + '=' * 80)
    print('PER-SECTION RETRIEVAL RATE (for source paper sections)')
    print('=' * 80)
    all_sections = set()
    for ev in [eval_orig, eval_v1, eval_v2]:
        all_sections.update(ev['per_section'].keys())

    # Show only common sections (>=20 attempts)
    common = []
    for sec in all_sections:
        attempts = max(eval_orig['per_section'].get(sec, {}).get('attempts', 0),
                       eval_v1['per_section'].get(sec, {}).get('attempts', 0),
                       eval_v2['per_section'].get(sec, {}).get('attempts', 0))
        if attempts >= 20:
            common.append((sec, attempts))
    common.sort(key=lambda x: -x[1])

    print(f'{"Section":<32} {"n":>6} {"Original":>10} {"V1":>10} {"V2":>10} {"Δ V2-O":>10}')
    print('-' * 80)
    for sec, n in common:
        o_rate = eval_orig['per_section'].get(sec, {}).get('rate', 0)
        v_rate = eval_v1['per_section'].get(sec, {}).get('rate', 0)
        s_rate = eval_v2['per_section'].get(sec, {}).get('rate', 0)
        d = (s_rate - o_rate) * 100
        print(f'  {sec:<30} {n:>6d} {o_rate:>9.2%} {v_rate:>9.2%} {s_rate:>9.2%}  {d:+9.2f} pp')

    # Save JSON
    out_json = HERE / 'retrieval_metrics.json'
    summary = {
        'original': {'aggregate': eval_orig['aggregate'],
                     'per_section': eval_orig['per_section'],
                     'n_queries': eval_orig['n_queries']},
        'v1_naive': {'aggregate': eval_v1['aggregate'],
                     'per_section': eval_v1['per_section'],
                     'n_queries': eval_v1['n_queries']},
        'v2_sh':    {'aggregate': eval_v2['aggregate'],
                     'per_section': eval_v2['per_section'],
                     'n_queries': eval_v2['n_queries']},
        'common_sections': [s for s, _ in common],
        'top_k': K,
    }
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(summary, f, indent=2)
    print(f'\n[Saved JSON] {out_json}')

    # Save Excel
    out_xlsx = HERE / 'retrieval_metrics_minimal.xlsx'
    build_excel(eval_orig, eval_v1, eval_v2, common, out_xlsx)
    print(f'[Saved Excel] {out_xlsx}')

    return summary


# ============================================================
# Excel builder (minimal)
# ============================================================
NAVY = '1F4E79'
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor=NAVY)
BODY_FONT = Font(name='Calibri', size=10)
BORDER = Border(left=Side(border_style='thin', color='CCCCCC'),
                right=Side(border_style='thin', color='CCCCCC'),
                top=Side(border_style='thin', color='CCCCCC'),
                bottom=Side(border_style='thin', color='CCCCCC'))


def build_excel(eval_orig, eval_v1, eval_v2, common_sections, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet 1: Aggregate summary
    ws = wb.create_sheet('1_Aggregate')
    ws['A1'] = 'Retrieval Metrics — 3-Way Comparison (Original / V1 Naive / V2 Schwartz-Hearst)'
    ws['A1'].font = Font(size=14, bold=True, color=NAVY)
    ws.merge_cells('A1:F1')

    headers = ['Metric', 'Original', 'V1 Naive', 'V2 SH', 'Δ V2-Orig (pp)', 'Δ V2-V1 (pp)']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', vertical='center')

    for row_i, metric in enumerate(['hit@5', 'precision@5', 'recall@5', 'mrr', 'ndcg@5', 'map@5'], 4):
        o = eval_orig['aggregate'][metric]
        v1 = eval_v1['aggregate'][metric]
        v2 = eval_v2['aggregate'][metric]
        cells = [metric, f'{o:.4f}', f'{v1:.4f}', f'{v2:.4f}',
                 f'{(v2-o)*100:+.2f}', f'{(v2-v1)*100:+.2f}']
        for col_i, val in enumerate(cells, 1):
            c = ws.cell(row=row_i, column=col_i, value=val)
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal='center' if col_i > 1 else 'left')
        # Color delta
        d_orig = (v2 - o) * 100
        ws.cell(row=row_i, column=5).fill = PatternFill('solid',
            fgColor='D1FAE5' if d_orig > 0.5 else ('FEE2E2' if d_orig < -0.5 else 'FFFFFF'))

    # Per-section
    ws['A12'] = 'Per-Section Retrieval Rate'
    ws['A12'].font = Font(size=12, bold=True, color=NAVY)

    headers2 = ['Section', 'n_attempts', 'Original', 'V1 Naive', 'V2 SH', 'Δ V2-Orig (pp)']
    for i, h in enumerate(headers2, 1):
        c = ws.cell(row=14, column=i, value=h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', vertical='center')

    row = 15
    for sec, n in common_sections:
        o = eval_orig['per_section'].get(sec, {}).get('rate', 0)
        v1 = eval_v1['per_section'].get(sec, {}).get('rate', 0)
        v2 = eval_v2['per_section'].get(sec, {}).get('rate', 0)
        cells = [sec, n, f'{o:.2%}', f'{v1:.2%}', f'{v2:.2%}', f'{(v2-o)*100:+.2f}']
        for col_i, val in enumerate(cells, 1):
            c = ws.cell(row=row, column=col_i, value=val)
            c.font = BODY_FONT; c.border = BORDER
            c.alignment = Alignment(horizontal='center' if col_i > 1 else 'left')
        d = (v2 - o) * 100
        ws.cell(row=row, column=6).fill = PatternFill('solid',
            fgColor='D1FAE5' if d > 1 else ('FEE2E2' if d < -1 else 'FFFFFF'))
        row += 1

    for col, w in [('A', 32), ('B', 12), ('C', 12), ('D', 12), ('E', 12), ('F', 16)]:
        ws.column_dimensions[col].width = w

    # Sheet 2: Per-question matrix
    ws2 = wb.create_sheet('2_Per_Question')
    ws2['A1'] = 'Per-Question Metrics (3 versions side-by-side)'
    ws2['A1'].font = Font(size=14, bold=True, color=NAVY)
    ws2.merge_cells('A1:N1')

    headers = ['idx', 'pubid', 'n_src',
               'Hit_O', 'Hit_V1', 'Hit_V2',
               'P5_O', 'P5_V1', 'P5_V2',
               'R5_O', 'R5_V1', 'R5_V2',
               'MRR_O', 'MRR_V1', 'MRR_V2',
               'nDCG_O', 'nDCG_V1', 'nDCG_V2']
    for i, h in enumerate(headers, 1):
        c = ws2.cell(row=3, column=i, value=h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', vertical='center')

    # Index by idx
    pq_orig = {m['idx']: m for m in eval_orig['per_query']}
    pq_v1 = {m['idx']: m for m in eval_v1['per_query']}
    pq_v2 = {m['idx']: m for m in eval_v2['per_query']}

    common_idx = sorted(set(pq_orig.keys()) & set(pq_v1.keys()) & set(pq_v2.keys()))
    for row_i, idx in enumerate(common_idx, 4):
        o, v1, v2 = pq_orig[idx], pq_v1[idx], pq_v2[idx]
        cells = [
            idx, o['source_pubid'], o['source_n_sections'],
            o['hit@5'], v1['hit@5'], v2['hit@5'],
            f"{o['precision@5']:.2f}", f"{v1['precision@5']:.2f}", f"{v2['precision@5']:.2f}",
            f"{o['recall@5']:.2f}", f"{v1['recall@5']:.2f}", f"{v2['recall@5']:.2f}",
            f"{o['mrr']:.3f}", f"{v1['mrr']:.3f}", f"{v2['mrr']:.3f}",
            f"{o['ndcg@5']:.3f}", f"{v1['ndcg@5']:.3f}", f"{v2['ndcg@5']:.3f}",
        ]
        for col_i, val in enumerate(cells, 1):
            c = ws2.cell(row=row_i, column=col_i, value=val)
            c.font = BODY_FONT
            c.alignment = Alignment(horizontal='center')

    ws2.freeze_panes = 'D4'
    ws2.auto_filter.ref = f'A3:R{3 + len(common_idx)}'

    wb.save(out_path)


if __name__ == '__main__':
    main()
