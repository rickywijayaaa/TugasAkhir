"""Build Excel for Schwartz-Hearst experiment results.

3-way comparison: Original baseline vs Naive expansion (V1) vs Schwartz-Hearst (V2).

Sheets:
  1_Overview            - 3-way summary stats + trajectory
  2_All_Wrong_V2        - 134 sampel salah di V2 (drilldown)
  3_NEW_Correct_V1_V2   - 15 sampel jadi correct V1->V2 (gain SH)
  4_NEW_Wrong_V1_V2     - 13 sampel jadi wrong V1->V2 (regresi SH)
  5_Total_Gain          - 29 sampel correct di V2 yang salah di Original
  6_Triple_Wrong        - 117 sampel salah di 3 versi (persistent failure)
  7_Per_Class           - per-kelas breakdown 3 versions
"""
import json
from pathlib import Path
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
RESULTS = HERE.parent.parent / 'results'
RESULTS_BM25 = RESULTS / 'BM25_Expansion'
OUT = HERE / 'comparison_3way_sh.xlsx'

NAVY = '1F4E79'; GOLD = 'D4A957'; GREEN = '4ADE80'; RED = 'F87171'
LRED = 'FEE2E2'; LGRN = 'D1FAE5'; LYEL = 'FEF3C7'; LBLU = 'DBEAFE'
LGRY = 'F3F4F6'; LPUR = 'E9D5FF'

HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor=NAVY)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
BODY_FONT = Font(name='Calibri', size=10)
WRAP = Alignment(vertical='top', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN = Side(border_style='thin', color='CCCCCC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row=1, last_col=None):
    if last_col is None: last_col = ws.max_column
    for col in range(1, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN; c.border = BORDER
    ws.row_dimensions[row].height = 32


def label_color(label):
    return {'yes': LGRN, 'no': LRED, 'maybe': LYEL}.get(label, '')


# ============================================================
# Load 3 datasets
# ============================================================
print('Loading 3 result sets...')
with open(RESULTS / 'baseline_openai_phase1_answers.json', encoding='utf-8') as f:
    orig = json.load(f)['results']
with open(RESULTS_BM25 / 'bm25expanded_baseline_openai_phase1_answers.json', encoding='utf-8') as f:
    v1 = json.load(f)['results']
with open(RESULTS_BM25 / 'sh_baseline_openai_phase1_answers.json', encoding='utf-8') as f:
    v2 = json.load(f)['results']

orig_idx = {r['idx']: r for r in orig}
v1_idx = {r['idx']: r for r in v1}
v2_idx = {r['idx']: r for r in v2}

# Categorize V1 -> V2 transitions
v1v2_new_correct, v1v2_new_wrong = [], []
v1v2_both_correct, v1v2_both_wrong = [], []
for idx in v1_idx:
    if idx not in v2_idx: continue
    o, n = v1_idx[idx], v2_idx[idx]
    if o['is_correct'] and n['is_correct']: v1v2_both_correct.append(idx)
    elif not o['is_correct'] and not n['is_correct']: v1v2_both_wrong.append(idx)
    elif not o['is_correct'] and n['is_correct']: v1v2_new_correct.append(idx)
    else: v1v2_new_wrong.append(idx)

# Categorize Original -> V2 transitions
ov2_new_correct, ov2_new_wrong = [], []
for idx in orig_idx:
    if idx not in v2_idx: continue
    o, n = orig_idx[idx], v2_idx[idx]
    if not o['is_correct'] and n['is_correct']: ov2_new_correct.append(idx)
    elif o['is_correct'] and not n['is_correct']: ov2_new_wrong.append(idx)

# Triple wrong
triple_wrong = [idx for idx in orig_idx if idx in v1_idx and idx in v2_idx
                and not orig_idx[idx]['is_correct']
                and not v1_idx[idx]['is_correct']
                and not v2_idx[idx]['is_correct']]

print(f'V1 -> V2 transitions: gain {len(v1v2_new_correct)}, loss {len(v1v2_new_wrong)}')
print(f'Original -> V2 total gain: {len(ov2_new_correct)}, loss: {len(ov2_new_wrong)}')
print(f'Triple wrong (persistent): {len(triple_wrong)}')


# ============================================================
# Sheet 1: Overview
# ============================================================
def build_overview(wb):
    ws = wb.create_sheet('1_Overview')
    ws['A1'] = 'SCHWARTZ-HEARST EXPERIMENT - 3-Way Comparison (n=500 PubMedQA)'
    ws['A1'].font = Font(size=15, bold=True, color=NAVY)
    ws.merge_cells('A1:H1')

    ws['A3'] = ('Comparing 3 versions of BM25 retrieval: '
                '(1) Original baseline (no expansion), '
                '(2) Naive expansion V1 (initial-based extractor), '
                '(3) Schwartz-Hearst V2 (handles internal-letter acronyms like HBO, VEGF, EGFR).')
    ws['A3'].font = Font(size=11, italic=True, color='6B7280')
    ws['A3'].alignment = WRAP
    ws.merge_cells('A3:H3')
    ws.row_dimensions[3].height = 48

    # Stats
    def compute(data):
        n = len(data)
        acc = sum(r['is_correct'] for r in data) / n
        accs = {}
        for lbl in ['yes', 'no', 'maybe']:
            sub = [r for r in data if r['ground_truth'] == lbl]
            accs[lbl] = sum(r['is_correct'] for r in sub) / len(sub) if sub else 0
        bal = sum(accs.values()) / 3
        return {'n': n, 'acc': acc, 'bal': bal, **{f'acc_{k}': v for k, v in accs.items()}}

    s_orig = compute(orig); s_v1 = compute(v1); s_v2 = compute(v2)

    headers = ['Metric', 'Original', 'Naive V1', 'SH V2', 'Δ V2-Orig', 'Δ V2-V1', 'Catatan']
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header(ws, row=5, last_col=len(headers))

    rows = [
        ('Overall accuracy',   s_orig['acc'],       s_v1['acc'],       s_v2['acc'],       ''),
        ('Balanced accuracy',  s_orig['bal'],       s_v1['bal'],       s_v2['bal'],       'Rata-rata 3 kelas'),
        ('Acc yes class',      s_orig['acc_yes'],   s_v1['acc_yes'],   s_v2['acc_yes'],   '275 sampel'),
        ('Acc no class',       s_orig['acc_no'],    s_v1['acc_no'],    s_v2['acc_no'],    '159 sampel'),
        ('Acc maybe class',    s_orig['acc_maybe'], s_v1['acc_maybe'], s_v2['acc_maybe'], '66 sampel (tersulit)'),
    ]
    row = 6
    for name, vo, vv1, vv2, note in rows:
        d_orig = (vv2 - vo) * 100
        d_v1 = (vv2 - vv1) * 100
        ws.cell(row=row, column=1, value=name).font = Font(size=11, bold=True, color=NAVY)
        ws.cell(row=row, column=2, value=f'{vo:.1%}')
        ws.cell(row=row, column=3, value=f'{vv1:.1%}')
        ws.cell(row=row, column=4, value=f'{vv2:.1%}')
        ws.cell(row=row, column=5, value=f'{d_orig:+.2f} pp')
        ws.cell(row=row, column=6, value=f'{d_v1:+.2f} pp')
        ws.cell(row=row, column=7, value=note)
        for c in [5, 6]:
            v_delta = d_orig if c == 5 else d_v1
            if v_delta > 0.5:
                ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=LGRN)
                ws.cell(row=row, column=c).font = Font(bold=True, color='065F46')
            elif v_delta < -0.5:
                ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=LRED)
                ws.cell(row=row, column=c).font = Font(bold=True, color='991B1B')
        for c in range(1, 8):
            ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=c).alignment = Alignment(
                vertical='center', horizontal='left' if c in (1, 7) else 'center')
        row += 1

    # Sample-level transitions
    row += 2
    ws.cell(row=row, column=1, value='V1 → V2 TRANSITIONS (Schwartz-Hearst impact)').font = Font(size=12, bold=True, color=NAVY)
    row += 1
    rows = [
        ('Both correct (no change)', len(v1v2_both_correct)),
        ('Both wrong (no change)',   len(v1v2_both_wrong)),
        ('NEW correct (V1 wrong, V2 correct)', len(v1v2_new_correct)),
        ('NEW wrong (V1 correct, V2 wrong)',   len(v1v2_new_wrong)),
    ]
    for name, count in rows:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=count)
        if 'NEW correct' in name:
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=LGRN)
        elif 'NEW wrong' in name:
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=LRED)
        ws.cell(row=row, column=2).alignment = CENTER
        ws.cell(row=row, column=2).font = Font(bold=True, size=12)
        row += 1
    net = len(v1v2_new_correct) - len(v1v2_new_wrong)
    ws.cell(row=row, column=1, value='Net V1 → V2').font = Font(size=12, bold=True, color=NAVY)
    ws.cell(row=row, column=2, value=f'+{net}' if net >= 0 else str(net))
    ws.cell(row=row, column=2).font = Font(size=14, bold=True, color='065F46' if net > 0 else '991B1B')
    ws.cell(row=row, column=2).alignment = CENTER

    # Total improvement
    row += 3
    ws.cell(row=row, column=1, value='TOTAL IMPROVEMENT (Original → V2)').font = Font(size=12, bold=True, color=NAVY)
    row += 1
    rows = [
        ('NEW correct (Orig wrong, V2 correct)', len(ov2_new_correct)),
        ('NEW wrong (Orig correct, V2 wrong)',   len(ov2_new_wrong)),
        ('Net Original → V2',                     len(ov2_new_correct) - len(ov2_new_wrong)),
        ('Triple wrong (persistent failure)',     len(triple_wrong)),
    ]
    for name, count in rows:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=f'+{count}' if 'Net' in name and count >= 0 else count)
        if 'NEW correct' in name or ('Net' in name and count > 0):
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=LGRN)
        elif 'NEW wrong' in name:
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=LRED)
        elif 'persistent' in name:
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=LYEL)
        ws.cell(row=row, column=2).alignment = CENTER
        ws.cell(row=row, column=2).font = Font(bold=True)
        row += 1

    # Idx 30 trajectory
    row += 2
    ws.cell(row=row, column=1, value='IDX 30 (HBO/NF) TRAJECTORY').font = Font(size=12, bold=True, color=NAVY)
    row += 1
    o30 = orig_idx.get(30); v1_30 = v1_idx.get(30); v2_30 = v2_idx.get(30)
    ws.cell(row=row, column=1, value='Question:').font = Font(bold=True)
    ws.cell(row=row, column=2, value=o30['question'])
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
    row += 1
    ws.cell(row=row, column=1, value='Ground truth:').font = Font(bold=True)
    ws.cell(row=row, column=2, value=o30['ground_truth'])
    ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=label_color(o30['ground_truth']))
    row += 1
    for ver, data in [('Original', o30), ('V1 Naive', v1_30), ('V2 Schwartz-Hearst', v2_30)]:
        ws.cell(row=row, column=1, value=ver)
        ws.cell(row=row, column=2, value=f"Pred: {data['predicted_label']}")
        ws.cell(row=row, column=3, value='CORRECT' if data['is_correct'] else 'WRONG')
        ws.cell(row=row, column=3).fill = PatternFill('solid',
            fgColor=LGRN if data['is_correct'] else LRED)
        ws.cell(row=row, column=3).font = Font(bold=True)
        ws.cell(row=row, column=3).alignment = CENTER
        row += 1

    # Navigation
    row += 2
    ws.cell(row=row, column=1, value='NAVIGASI SHEET').font = Font(size=12, bold=True, color=NAVY)
    row += 1
    nav = [
        '→ Sheet "2_All_Wrong_V2"      : 134 sampel masih salah di V2 (drilldown semua)',
        '→ Sheet "3_NEW_Correct_V1_V2" : 15 sampel gain dari V1 ke V2 (efek SH)',
        '→ Sheet "4_NEW_Wrong_V1_V2"   : 13 sampel regresi V1 ke V2 (investigate)',
        '→ Sheet "5_Total_Gain"        : 29 sampel total improvement Orig ke V2',
        '→ Sheet "6_Triple_Wrong"      : 117 sampel persistent failure (semua 3 versi salah)',
        '→ Sheet "7_Per_Class"         : per-kelas yes/no/maybe breakdown',
    ]
    for n in nav:
        ws.cell(row=row, column=1, value=n).font = Font(size=11)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1

    widths = {'A': 32, 'B': 18, 'C': 18, 'D': 18, 'E': 14, 'F': 14, 'G': 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ============================================================
# Generic 3-way diff sheet
# ============================================================
def build_3way_sheet(wb, sheet_name, idx_list, title, intro):
    ws = wb.create_sheet(sheet_name)
    ws['A1'] = title
    ws['A1'].font = Font(size=14, bold=True, color=NAVY)
    ws.merge_cells('A1:M1')

    ws['A2'] = intro
    ws['A2'].font = Font(size=10, italic=True, color='6B7280')
    ws['A2'].alignment = WRAP
    ws.merge_cells('A2:M2')
    ws.row_dimensions[2].height = 40

    headers = ['Idx', 'PMID', 'GT', 'Orig Pred', 'V1 Pred', 'V2 Pred',
               'Question', 'V1: Source sections', 'V2: Source sections',
               'V2 Answer (preview)', 'Reference', 'Section labels (V2)', 'V2 RRF top scores']
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, row=4, last_col=len(headers))

    for row_i, idx in enumerate(idx_list, 5):
        o = orig_idx.get(idx); v1r = v1_idx.get(idx); v2r = v2_idx.get(idx)
        if not (o and v1r and v2r): continue
        target_pubid = str(o['pubid'])

        v1_secs = []
        if 'context_pubids' in v1r:
            for pid, sec in zip(v1r['context_pubids'], v1r.get('context_sections', ['?']*5)):
                if pid == target_pubid:
                    v1_secs.append(sec)

        v2_secs = []
        for pid, sec in zip(v2r['context_pubids'], v2r['context_sections']):
            if pid == target_pubid:
                v2_secs.append(sec)

        cells = [
            idx, target_pubid, o['ground_truth'],
            o['predicted_label'], v1r['predicted_label'], v2r['predicted_label'],
            o['question'],
            ', '.join(v1_secs) if v1_secs else '(none)',
            ', '.join(v2_secs) if v2_secs else '(none)',
            v2r['answer'][:300] + ('...' if len(v2r['answer']) > 300 else ''),
            o.get('reference', '-'),
            ', '.join(v2r.get('context_sections', [])),
            ', '.join(f'{s:.3f}' for s in v2r.get('rrf_scores', [])[:5]),
        ]
        for col_i, v in enumerate(cells, 1):
            c = ws.cell(row=row_i, column=col_i, value=v)
            c.font = BODY_FONT; c.alignment = WRAP; c.border = BORDER

        ws.cell(row=row_i, column=2).hyperlink = f'https://pubmed.ncbi.nlm.nih.gov/{target_pubid}/'
        ws.cell(row=row_i, column=2).font = Font(color=NAVY, underline='single', size=10)

        ws.cell(row=row_i, column=3).fill = PatternFill('solid', fgColor=label_color(o['ground_truth']))
        for col_i, ver_data in enumerate([o, v1r, v2r], 4):
            ws.cell(row=row_i, column=col_i).fill = PatternFill('solid',
                fgColor=LGRN if ver_data['predicted_label'] == o['ground_truth'] else LRED)
            ws.cell(row=row_i, column=col_i).alignment = CENTER

        # Highlight if V2 retrieves more source sections than V1
        if len(v2_secs) > len(v1_secs) and v2_secs:
            ws.cell(row=row_i, column=9).fill = PatternFill('solid', fgColor=LGRN)
            ws.cell(row=row_i, column=9).font = Font(bold=True, color='065F46', size=10)

    widths = [6, 12, 6, 9, 9, 9, 38, 25, 25, 50, 35, 30, 25]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'G5'
    if idx_list:
        ws.auto_filter.ref = f'A4:M{4 + len(idx_list)}'


# ============================================================
# Per-class stats (3-way)
# ============================================================
def build_per_class(wb):
    ws = wb.create_sheet('7_Per_Class')
    ws['A1'] = 'Per-Class Accuracy: 3-Way Comparison'
    ws['A1'].font = Font(size=14, bold=True, color=NAVY)
    ws.merge_cells('A1:I1')

    ws['A3'] = ('Per-kelas accuracy untuk 3 versi. Lihat di kelas mana SH V2 paling membantu '
                'dan di mana terjadi regresi vs V1.')
    ws['A3'].font = Font(size=10, italic=True, color='6B7280')
    ws['A3'].alignment = WRAP
    ws.merge_cells('A3:I3')

    headers = ['Kelas GT', 'n_total', 'Orig acc', 'V1 acc', 'V2 acc',
               'Δ V1-Orig (pp)', 'Δ V2-V1 (pp)', 'Δ V2-Orig (pp)', 'Net (V2-Orig)']
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header(ws, row=5, last_col=len(headers))

    row = 6
    for lbl in ['yes', 'no', 'maybe']:
        sub_o = [r for r in orig if r['ground_truth'] == lbl]
        sub_1 = [r for r in v1 if r['ground_truth'] == lbl]
        sub_2 = [r for r in v2 if r['ground_truth'] == lbl]
        n = len(sub_o)
        c_o = sum(r['is_correct'] for r in sub_o)
        c_1 = sum(r['is_correct'] for r in sub_1)
        c_2 = sum(r['is_correct'] for r in sub_2)
        d1 = (c_1 - c_o) / n * 100
        d2 = (c_2 - c_1) / n * 100
        d_total = (c_2 - c_o) / n * 100

        cells = [
            lbl, n,
            f'{c_o/n:.1%}', f'{c_1/n:.1%}', f'{c_2/n:.1%}',
            f'{d1:+.2f}', f'{d2:+.2f}', f'{d_total:+.2f}',
            f'{c_2-c_o:+d}'
        ]
        for col_i, v in enumerate(cells, 1):
            c = ws.cell(row=row, column=col_i, value=v)
            c.alignment = CENTER; c.border = BORDER
            if col_i == 1:
                c.fill = PatternFill('solid', fgColor=label_color(lbl))
                c.font = Font(bold=True)
            else:
                c.font = Font(size=10)

        # Color delta columns
        for col_i, delta in [(6, d1), (7, d2), (8, d_total)]:
            if delta > 1:
                ws.cell(row=row, column=col_i).fill = PatternFill('solid', fgColor=LGRN)
                ws.cell(row=row, column=col_i).font = Font(bold=True, color='065F46')
            elif delta < -1:
                ws.cell(row=row, column=col_i).fill = PatternFill('solid', fgColor=LRED)
                ws.cell(row=row, column=col_i).font = Font(bold=True, color='991B1B')
        row += 1

    # Insight
    row += 2
    ws.cell(row=row, column=1, value='INSIGHT KUNCI').font = Font(size=12, bold=True, color=NAVY)
    row += 1
    insights = [
        '• Yes class: V2 paling tinggi (90.2%) — SH menambah konteks yang tepat untuk positive findings.',
        '• No class: V1 dan V2 sama (72.3%, +9.4 pp dari Original) — gain utama datang dari naive expansion.',
        '• Maybe class: V2 lebih rendah dari V1 (4.5% vs 7.6%) — REGRESI! Mungkin karena context lebih lengkap membuat LLM commit ke yes/no.',
        '• SH manfaat terutama untuk kelas yes (lebih banyak context yang relevan)',
        '• Trade-off: lebih banyak konteks bagus → lebih sedikit "maybe" prediction (commitment effect)',
    ]
    for ins in insights:
        ws.cell(row=row, column=1, value=ins).font = Font(size=11)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        ws.cell(row=row, column=1).alignment = WRAP
        row += 1

    widths = [12, 10, 12, 12, 12, 14, 14, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================
# MAIN
# ============================================================
if __name__ == '__main__':
    wb = Workbook()
    wb.remove(wb.active)

    print('Building Sheet 1: Overview...')
    build_overview(wb)

    all_wrong_v2 = [idx for idx in v2_idx if not v2_idx[idx]['is_correct']]
    print(f'Building Sheet 2: All Wrong V2 ({len(all_wrong_v2)} samples)...')
    build_3way_sheet(wb, '2_All_Wrong_V2', all_wrong_v2,
                     f'Semua sampel SALAH di Schwartz-Hearst V2 ({len(all_wrong_v2)}/500)',
                     'Drilldown semua kasus yang masih salah setelah Schwartz-Hearst expansion. '
                     'Filter via auto-filter: lihat distribusi GT, bandingkan prediksi 3 versi, '
                     'cek sections paper sumber yang ter-retrieve di V2 vs V1.')

    print(f'Building Sheet 3: NEW Correct V1->V2 ({len(v1v2_new_correct)} samples)...')
    build_3way_sheet(wb, '3_NEW_Correct_V1_V2', v1v2_new_correct,
                     f'NEW CORRECT V1→V2: {len(v1v2_new_correct)} sampel (efek murni Schwartz-Hearst)',
                     'Sampel di mana naive V1 SALAH tapi SH V2 BENAR. '
                     'Kolom "V2: Source sections" highlight hijau = lebih banyak section paper sumber '
                     'ter-retrieve berkat SH (mis. RESULTS yang sebelumnya miss).')

    print(f'Building Sheet 4: NEW Wrong V1->V2 ({len(v1v2_new_wrong)} samples)...')
    build_3way_sheet(wb, '4_NEW_Wrong_V1_V2', v1v2_new_wrong,
                     f'NEW WRONG V1→V2 (regresi): {len(v1v2_new_wrong)} sampel',
                     'Investigate: apakah SH expansion menambah noise atau over-expand? '
                     'Cek apakah SH retrieve sections berbeda yang membingungkan LLM.')

    print(f'Building Sheet 5: Total Gain Original->V2 ({len(ov2_new_correct)} samples)...')
    build_3way_sheet(wb, '5_Total_Gain', ov2_new_correct,
                     f'TOTAL GAIN Original→V2: {len(ov2_new_correct)} sampel (story untuk laporan)',
                     'Sampel yang fail di Original baseline tapi correct di V2 SH. '
                     'Ini story improvement utama untuk laporan TA: gain dari semua intervention '
                     '(naive + SH expansion).')

    print(f'Building Sheet 6: Triple Wrong ({len(triple_wrong)} samples)...')
    build_3way_sheet(wb, '6_Triple_Wrong', triple_wrong,
                     f'PERSISTENT FAILURE: {len(triple_wrong)} sampel salah di SEMUA 3 versi',
                     'Bukan masalah expansion. Kemungkinan: information gap struktural '
                     '(jawaban di CONCLUSIONS yang tidak di-retrieve), reasoning failure LLM, '
                     'atau prompt anti-maybe bias yang mendominasi. Kandidat untuk future work.')

    print('Building Sheet 7: Per-Class Stats...')
    build_per_class(wb)

    wb.save(OUT)
    print(f'\n[Saved] {OUT}')
    print(f'  Size: {OUT.stat().st_size / 1024:.1f} KB')
